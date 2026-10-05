import io
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def create_presentation_pptx(data: dict) -> io.BytesIO:
    """Формирует файл презентации PowerPoint из JSON данных"""
    prs = Presentation()

    # Устанавливаем современный формат слайдов 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- СЛАЙД 1: Титульный ---
    title_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_layout)

    title = title_slide.shapes.title
    subtitle = title_slide.placeholders[1]

    title.text = data.get("title", "Презентация к уроку")
    subtitle.text = f"Предмет: {data.get('subject', 'Литература')} | Класс: {data.get('grade', '')}"

    # --- СЛАЙДЫ 2..N: Содержимое ---
    content_layout = prs.slide_layouts[1]  # Шаблон: Заголовок + Текст

    for s_data in data.get("slides", []):
        slide = prs.slides.add_slide(content_layout)

        # Заголовок слайда
        slide_title = slide.shapes.title
        slide_title.text = (
            f"{s_data.get('slide_number', '')}. {s_data.get('title', '')}"
        )

        # Основные тезисы (Bullet Points)
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        tf.clear()

        bullets = s_data.get("content_bullets", [])
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(20)
            p.space_after = Pt(12)

        # Заметки докладчика (Заметки для учителя снизу слайда)
        if s_data.get("teacher_notes"):
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            text_frame.text = (
                f"Подсказка учителя: {s_data.get('teacher_notes')}"
            )

    file_stream = io.BytesIO()
    prs.save(file_stream)
    file_stream.seek(0)
    return file_stream