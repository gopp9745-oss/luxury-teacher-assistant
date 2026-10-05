import io
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def create_test_docx(data: dict, doc_mode: str = "full") -> io.BytesIO:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_text = data.get("title", "Проверочная работа")
    subject = data.get("subject", "Предмет")
    grade = data.get("grade", "Класс")
    questions = data.get("questions", [])

    # === РАЗДЕЛ ДЛЯ УЧЕНИКОВ ===
    if doc_mode in ["student", "full"]:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(title_text)
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(18)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_meta = p_meta.add_run(f"Предмет: {subject} | {grade}")
        run_meta.font.name = "Calibri"
        run_meta.font.size = Pt(11)
        run_meta.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        p_info = doc.add_paragraph()
        p_info.paragraph_format.space_after = Pt(18)
        run_info = p_info.add_run(
            "ФИО ученика: ____________________________________   Дата:"
            " ____________"
        )
        run_info.font.name = "Calibri"
        run_info.font.size = Pt(11)

        for q in questions:
            q_num = q.get("id", "")
            q_text = q.get("question", "")
            q_type = q.get("type", "single")

            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(12)
            p_q.paragraph_format.space_after = Pt(4)

            r_num = p_q.add_run(f"Задание {q_num}. ")
            r_num.font.bold = True
            r_num.font.size = Pt(12)

            r_q = p_q.add_run(q_text)
            r_q.font.size = Pt(12)

            if q_type == "matching":
                left = q.get("left_column", [])
                right = q.get("right_column", [])

                table = doc.add_table(rows=1, cols=2)
                table.autofit = False

                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = "Левая колонка"
                hdr_cells[1].text = "Правая колонка"
                hdr_cells[0].paragraphs[0].runs[0].font.bold = True
                hdr_cells[1].paragraphs[0].runs[0].font.bold = True

                max_len = max(len(left), len(right))
                for i in range(max_len):
                    row_cells = table.add_row().cells
                    row_cells[0].text = left[i] if i < len(left) else ""
                    row_cells[1].text = right[i] if i < len(right) else ""

                p_ans_place = doc.add_paragraph()
                p_ans_place.paragraph_format.space_before = Pt(6)
                p_ans_place.add_run(
                    "Ответ (запишите пары): ______________________"
                ).font.italic = True
            else:
                for opt in q.get("options", []):
                    p_opt = doc.add_paragraph()
                    p_opt.paragraph_format.left_indent = Inches(0.25)
                    p_opt.paragraph_format.space_after = Pt(2)
                    r_opt = p_opt.add_run(f"[  ] {opt}")
                    r_opt.font.size = Pt(11)

    if doc_mode == "full":
        doc.add_page_break()

    # === РАЗДЕЛ ДЛЯ УЧИТЕЛЯ ===
    if doc_mode in ["teacher", "full"]:
        p_t_title = doc.add_paragraph()
        p_t_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t_title = p_t_title.add_run("КЛЮЧИ И ПОЯСНЕНИЯ ДЛЯ УЧИТЕЛЯ")
        r_t_title.font.name = "Calibri"
        r_t_title.font.size = Pt(16)
        r_t_title.font.bold = True
        r_t_title.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)

        p_t_meta = doc.add_paragraph()
        p_t_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t_meta.paragraph_format.space_after = Pt(18)
        r_t_meta = p_t_meta.add_run(f"К работе: {title_text} ({subject}, {grade})")
        r_t_meta.font.size = Pt(11)

        for q in questions:
            q_num = q.get("id", "")
            q_text = q.get("question", "")
            q_type = q.get("type", "single")
            exp = q.get("explanation", "")

            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.space_before = Pt(10)
            p_ans.paragraph_format.space_after = Pt(2)

            r_ans_num = p_ans.add_run(f"Задание {q_num}: ")
            r_ans_num.font.bold = True
            p_ans.add_run(q_text).font.italic = True

            p_corr = doc.add_paragraph()
            p_corr.paragraph_format.left_indent = Inches(0.2)
            p_corr.paragraph_format.space_after = Pt(2)

            r_c_label = p_corr.add_run("Правильный ответ: ")
            r_c_label.font.bold = True
            r_c_label.font.color.rgb = RGBColor(0x05, 0x96, 0x69)

            if q_type == "matching":
                pairs = ", ".join(q.get("correct_pairs", []))
                p_corr.add_run(pairs).font.bold = True
            else:
                correct_list = q.get("correct_answers", [])
                if not correct_list and "correct_answer" in q:
                    correct_list = [q["correct_answer"]]
                p_corr.add_run(", ".join(correct_list)).font.bold = True

            if exp:
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.left_indent = Inches(0.2)
                p_exp.paragraph_format.space_after = Pt(8)
                p_exp.add_run("Пояснение: ").font.bold = True
                r_e_val = p_exp.add_run(exp)
                r_e_val.font.size = Pt(10)
                r_e_val.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream